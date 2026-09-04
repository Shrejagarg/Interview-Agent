import os
import re
import logging
from .config import get_config
from .skills_database import ALL_SKILLS, SKILL_CATEGORIES, MARKETING_SKILLS, SKILLS_BY_DOMAIN

logger = logging.getLogger(__name__)
cfg = get_config()

EXPERIENCE_PATTERNS = [
    r"(\d+)\+?\s*years?\s*(?:of\s+)?(?:experience|exp)",
    r"experience\s*[:\-]?\s*(\d+)\+?\s*years?",
    r"(\d+)\+?\s*yrs?\s*(?:of\s+)?(?:experience|exp)",
]

EDUCATION_KEYWORDS = [
    "bachelor", "master", "phd", "mba", "b.sc", "b.s.", "m.sc", "m.s.",
    "b.tech", "b.tech.", "m.tech", "m.tech.", "bca", "mca", "bba", "mba",
    "degree", "university", "college", "institute", "graduated", "diploma"
]

JOB_TITLE_KEYWORDS = [
    "manager", "director", "lead", "head", "chief", "officer", "specialist",
    "executive", "coordinator", "analyst", "consultant", "associate",
    "intern", "trainee", "fresher", "junior", "senior", "vp", "cmo",
    "marketing manager", "marketing executive", "marketing specialist",
    "brand manager", "content manager", "social media manager",
    "digital marketing manager", "growth manager", "product manager"
]

SENIOR_KEYWORDS = ["director", "vp", "chief", "head", "lead", "principal", "senior"]
MID_KEYWORDS = ["manager", "specialist", "consultant", "analyst", "coordinator"]
FRESHER_KEYWORDS = ["intern", "trainee", "fresher", "junior", "assistant", "associate"]

SECTION_HEADERS = [
    "summary", "profile", "objective", "experience", "work experience",
    "employment", "education", "skills", "technical skills", "certifications",
    "projects", "achievements", "awards", "publications", "languages",
    "interests", "hobbies", "references", "contact"
]


def extract_text_from_pdf(file_path):
    import pdfplumber
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        logger.info("Extracted %d chars from PDF: %s", len(text), file_path)
    except Exception as e:
        logger.error("Failed to read PDF %s: %s", file_path, e)
    return text


def extract_text_from_docx(file_path):
    from docx import Document
    text = ""
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"
        logger.info("Extracted %d chars from DOCX: %s", len(text), file_path)
    except Exception as e:
        logger.error("Failed to read DOCX %s: %s", file_path, e)
    return text


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    supported = cfg["resume"]["supported_formats"]

    if ext not in supported:
        logger.error("Unsupported format '%s'. Supported: %s", ext, supported)
        return ""

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    elif ext == ".txt":
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1") as f:
                return f.read()
    return ""


def extract_contact(text):
    contact = {"email": None, "phone": None}

    email_match = re.search(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", text)
    if email_match:
        contact["email"] = email_match.group(0)

    phone_match = re.search(r"[\+]?[\d\-\(\)\s]{7,15}", text)
    if phone_match:
        raw = phone_match.group(0).strip()
        digits = re.sub(r"\D", "", raw)
        if 7 <= len(digits) <= 15:
            contact["phone"] = raw.strip()

    return contact


def _is_header_or_noise(line):
    line_lower = line.lower().strip()
    if len(line_lower) <= 2:
        return True
    if re.match(r"^[\+\d\-\(\)\s]{7,}$", line):
        return True
    if re.search(r"[@]|\.com|\.in|\.org|\.net", line, re.I):
        return True
    if any(line_lower == h for h in SECTION_HEADERS):
        return True
    if re.match(r"^[\-_=|:/\\]+$", line):
        return True
    return False


def _looks_like_name(line):
    words = line.split()
    if len(words) < 1 or len(words) > 4:
        return False
    if any(kw in line.lower() for kw in JOB_TITLE_KEYWORDS):
        return False
    if any(kw in line.lower() for kw in SECTION_HEADERS):
        return False
    if re.search(r"[@]|\.com|\.in|\.org|\.net", line, re.I):
        return False
    if re.match(r"^[\+\d\-\(\)\s]{7,}$", line):
        return False
    if re.match(r"^[\-_=|:/\\]+$", line):
        return False
    if all(c.isalpha() or c.isspace() for c in line):
        return True
    return False


def extract_name(text):
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    if not lines:
        return "Unknown"

    for line in lines:
        if _is_header_or_noise(line):
            continue
        if _looks_like_name(line):
            return line
        break

    return "Unknown"


def extract_skills(text):
    """Extract skills from resume text.

    Returns per-domain skill matches (per_domain_skills) in addition to the
    flat list used by the existing marketing pipeline.
    """
    text_lower = text.lower()
    found: list[str] = []
    found_categories: set[str] = set()
    per_domain: dict[str, list[str]] = {domain: [] for domain in SKILLS_BY_DOMAIN}

    for skill in ALL_SKILLS:
        pattern = r"\b" + re.escape(skill) + r"\b"
        if re.search(pattern, text_lower):
            found.append(skill)
            cat = SKILL_CATEGORIES.get(skill, "unknown")
            found_categories.add(cat)
            # Attribute skill to every domain that contains it
            for domain, cats in SKILLS_BY_DOMAIN.items():
                for cat_skills in cats.values():
                    if skill in cat_skills:
                        per_domain[domain].append(skill)

    # De-duplicate per-domain lists
    per_domain = {d: sorted(set(skills)) for d, skills in per_domain.items()}

    return {
        "skills":            sorted(set(found)),
        "categories":        sorted(found_categories),
        "skill_count":       len(set(found)),
        "per_domain_skills": per_domain,
    }


# Domain keyword signals (titles + domain-specific jargon)
_DOMAIN_TITLE_SIGNALS: dict[str, list[str]] = {
    "marketing":            ["marketing", "brand", "content", "seo", "social media", "growth", "cmo", "demand gen"],
    "software_engineering": ["engineer", "developer", "programmer", "software", "backend", "frontend", "devops", "architect", "cto"],
    "finance":              ["finance", "analyst", "investment", "banker", "accounting", "cfo", "portfolio", "trading", "equity"],
    "hr":                   ["human resources", "hr", "recruiter", "talent", "people operations", "hrbp", "chro"],
    "sales":                ["sales", "account executive", "business development", "bdr", "sdr", "revenue", "closing"],
}


def infer_domains(resume_data: dict) -> list[str]:
    """Score each domain and return a ranked list (most likely first).

    Scoring:
      - +2 per matched per-domain skill
      - +3 per matched title keyword
    """
    scores: dict[str, float] = {domain: 0.0 for domain in SKILLS_BY_DOMAIN}

    # Skill-based signals
    per_domain = resume_data.get("per_domain_skills") or {}
    for domain, skills in per_domain.items():
        scores[domain] += len(skills) * 2

    # Title-based signals
    titles_text = " ".join(resume_data.get("job_titles", [])).lower()
    for domain, signals in _DOMAIN_TITLE_SIGNALS.items():
        for sig in signals:
            if sig in titles_text:
                scores[domain] += 3

    ranked = sorted(scores, key=lambda d: scores[d], reverse=True)
    return [d for d in ranked if scores[d] > 0]


def extract_experience(text):
    text_lower = text.lower()
    years = None

    for pattern in EXPERIENCE_PATTERNS:
        match = re.search(pattern, text_lower)
        if match:
            years = int(match.group(1))
            break

    job_titles = []
    for line in text.split("\n"):
        line_lower = line.lower().strip()
        for title in JOB_TITLE_KEYWORDS:
            if title in line_lower:
                job_titles.append(line.strip())
                break

    return {
        "years": years,
        "job_titles": list(set(job_titles))[:10],
        "title_count": len(set(job_titles))
    }


def extract_education(text):
    text_lower = text.lower()
    education = []

    for line in text.split("\n"):
        line_lower = line.lower().strip()
        for keyword in EDUCATION_KEYWORDS:
            if keyword in line_lower:
                education.append(line.strip())
                break

    return {
        "entries": list(set(education))[:5],
        "has_education": len(education) > 0
    }


def detect_experience_level(experience_data):
    years = experience_data.get("years")
    titles = experience_data.get("job_titles", [])
    titles_lower = " ".join(titles).lower()

    if years is not None:
        if years <= 1:
            return "fresher"
        elif years <= 4:
            return "mid"
        else:
            return "senior"

    for kw in SENIOR_KEYWORDS:
        if kw in titles_lower:
            return "senior"
    for kw in MID_KEYWORDS:
        if kw in titles_lower:
            return "mid"
    for kw in FRESHER_KEYWORDS:
        if kw in titles_lower:
            return "fresher"

    return "unknown"


def score_resume_quality(resume_data):
    w = cfg["resume"]["quality_weights"]
    thresholds = w["skill_thresholds"]
    length_thresholds = w["length_thresholds"]

    score = 0
    breakdown = {}

    if resume_data.get("contact", {}).get("email"):
        score += w["contact_info"]
        breakdown["contact_info"] = w["contact_info"]
    else:
        breakdown["contact_info"] = 0

    skill_count = resume_data.get("skills", {}).get("skill_count", 0)
    if skill_count >= thresholds["high"]:
        score += w["skills_high"]
        breakdown["skills"] = w["skills_high"]
    elif skill_count >= thresholds["mid"]:
        score += w["skills_mid"]
        breakdown["skills"] = w["skills_mid"]
    elif skill_count >= 1:
        score += w["skills_low"]
        breakdown["skills"] = w["skills_low"]
    else:
        breakdown["skills"] = 0

    exp = resume_data.get("experience", {})
    if exp.get("years") is not None:
        score += w["experience_years"]
        breakdown["experience_declared"] = w["experience_years"]
    elif exp.get("title_count", 0) > 0:
        score += w["experience_titles"]
        breakdown["experience_declared"] = w["experience_titles"]
    else:
        breakdown["experience_declared"] = 0

    edu = resume_data.get("education", {})
    if edu.get("has_education"):
        score += w["education"]
        breakdown["education"] = w["education"]
    else:
        breakdown["education"] = 0

    name = resume_data.get("name", "Unknown")
    if name != "Unknown":
        score += w["name"]
        breakdown["name"] = w["name"]
    else:
        breakdown["name"] = 0

    raw_text = resume_data.get("raw_text", "")
    word_count = len(raw_text.split())
    if word_count >= length_thresholds["high"]:
        score += w["length_high"]
        breakdown["length"] = w["length_high"]
    elif word_count >= length_thresholds["mid"]:
        score += w["length_mid"]
        breakdown["length"] = w["length_mid"]
    else:
        breakdown["length"] = 0

    return {
        "score": min(score, 100),
        "max_score": 100,
        "breakdown": breakdown,
        "word_count": word_count
    }


def check_required_sections(resume_data):
    required = cfg["resume"]["required_sections"]
    missing = []

    for section in required:
        if section == "contact":
            if not resume_data.get("contact", {}).get("email"):
                missing.append("contact")
        elif section == "skills":
            if resume_data.get("skills", {}).get("skill_count", 0) == 0:
                missing.append("skills")
        elif section == "experience":
            exp = resume_data.get("experience", {})
            if exp.get("years") is None and exp.get("title_count", 0) == 0:
                missing.append("experience")

    return {
        "all_present": len(missing) == 0,
        "missing": missing,
        "required": required
    }


def parse_resume(file_path):
    if not os.path.exists(file_path):
        logger.error("Resume file not found: %s", file_path)
        return None

    ext = os.path.splitext(file_path)[1].lower()
    supported = cfg["resume"]["supported_formats"]
    if ext not in supported:
        logger.error("Unsupported format '%s'. Supported: %s", ext, supported)
        return None

    raw_text = extract_text(file_path)
    if not raw_text.strip():
        logger.error("No text extracted from: %s", file_path)
        return None

    contact = extract_contact(raw_text)
    name = extract_name(raw_text)
    skills = extract_skills(raw_text)
    experience = extract_experience(raw_text)
    education = extract_education(raw_text)
    experience_level = detect_experience_level(experience)

    resume_data = {
        "file_path":        file_path,
        "name":             name,
        "contact":          contact,
        "skills":           skills,
        "per_domain_skills": skills.get("per_domain_skills", {}),
        "experience":       experience,
        "education":        education,
        "experience_level": experience_level,
        "raw_text":         raw_text,
        "job_titles":       experience.get("job_titles", []),
    }

    # Infer which domains the resume aligns with (ranked list)
    resume_data["matched_domains"] = infer_domains(resume_data)

    resume_data["quality"] = score_resume_quality(resume_data)
    resume_data["sections_check"] = check_required_sections(resume_data)

    quality_score = resume_data["quality"]["score"]
    min_quality = cfg["resume"]["min_quality_score"]
    resume_data["quality"]["meets_minimum"] = quality_score >= min_quality

    logger.info(
        "Parsed resume: %s (%d skills, %s level, quality %d/100, top domains: %s)",
        name, skills["skill_count"], experience_level,
        quality_score, resume_data["matched_domains"][:3],
    )

    return resume_data

